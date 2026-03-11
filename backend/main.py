"""
NS Checker Web App — FastAPI backend.
Handles CSV upload, concurrent DNS NS lookups, SSE progress, and Excel download.
"""

import csv
import io
import json
import time
import uuid
import asyncio
import threading
from pathlib import Path
from typing import AsyncGenerator

import pandas as pd
from fastapi import FastAPI, File, UploadFile, HTTPException
from fastapi.responses import StreamingResponse, FileResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from openpyxl.styles import Font, PatternFill, Alignment

from .ns_checker import run_bulk_lookup

app = FastAPI(title="NS Checker")

# In-memory job store: job_id -> {status, progress, total, results, error}
jobs: dict[str, dict] = {}
RESULTS_DIR = Path("/tmp/ns_results")
RESULTS_DIR.mkdir(exist_ok=True)


def parse_csv_domains(content: bytes) -> list[str]:
    """Extract domains from column A of a CSV, skipping obvious headers."""
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    domains = []
    for i, row in enumerate(reader):
        if not row:
            continue
        val = row[0].strip()
        if not val:
            continue
        if i == 0 and val.lower() in {"domain", "domains", "url", "website", "hostname", "host"}:
            continue
        domains.append(val)
    return domains


def extract_provider(nameservers: list[str]) -> str:
    """Return the unique NS base-domain(s) from a list of nameserver hostnames.
    e.g. ['ns41.dnsmadeeasy.com', 'ns11.dnsmadeeasy.com'] -> 'dnsmadeeasy.com'
    """
    providers = set()
    for ns in nameservers:
        parts = ns.rstrip(".").lower().split(".")
        if len(parts) >= 2:
            providers.add(".".join(parts[-2:]))
    return ", ".join(sorted(providers)) if providers else "—"


def _to_rows(results: list[dict]) -> list[dict]:
    return [
        {
            "Domain":      r["domain"],
            "Nameservers": ", ".join(r["nameservers"]),
            "Provider":    r.get("provider") or extract_provider(r["nameservers"]),
            "NS Count":    r["ns_count"],
            "Status":      r["status"],
        }
        for r in results
    ]


def _style_sheet(ws) -> None:
    """Apply header styling and status colour-coding to a worksheet."""
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    ok_fill   = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    err_fill  = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    warn_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    star_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")  # gold tint for flagged sheet

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        status = str(row[4].value or "")  # column E
        if status == "OK":
            row[4].fill = ok_fill
        elif status in ("NXDOMAIN", "TIMEOUT") or status.startswith("ERROR"):
            row[4].fill = err_fill
        else:
            row[4].fill = warn_fill

    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 80)


def build_excel(results: list[dict], starred_providers: list[str] | None = None) -> bytes:
    """
    Generate a formatted Excel workbook.
    If starred_providers is supplied, produce two sheets:
      - "Flagged for Migration"  — domains whose provider is in starred_providers
      - "Other Domains"          — everything else
    Otherwise a single "NS Results" sheet.
    """
    buf = io.BytesIO()
    starred_set = {p.lower() for p in (starred_providers or [])}

    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        if starred_set:
            flagged = [r for r in results if any(
                p.strip().lower() in starred_set for p in (r.get("provider") or "").split(", ")
            )]
            other   = [r for r in results if r not in flagged]

            pd.DataFrame(_to_rows(flagged)).to_excel(writer, index=False, sheet_name="Flagged for Migration")
            pd.DataFrame(_to_rows(other)).to_excel(writer,   index=False, sheet_name="Other Domains")

            _style_sheet(writer.sheets["Flagged for Migration"])
            _style_sheet(writer.sheets["Other Domains"])

            # Extra gold header tint on the flagged sheet to make it obvious
            star_header = PatternFill(start_color="BF8F00", end_color="BF8F00", fill_type="solid")
            for cell in writer.sheets["Flagged for Migration"][1]:
                cell.fill = star_header
        else:
            pd.DataFrame(_to_rows(results)).to_excel(writer, index=False, sheet_name="NS Results")
            _style_sheet(writer.sheets["NS Results"])

    buf.seek(0)
    return buf.read()


def run_job(job_id: str, domains: list[str]) -> None:
    """Run in a background thread: perform lookups, update job state."""
    job = jobs[job_id]
    job["status"] = "running"
    job["total"] = len(domains)
    job["progress"] = 0

    def on_progress(completed: int, total: int) -> None:
        job["progress"] = completed

    try:
        results = run_bulk_lookup(domains, workers=20, timeout=5.0, progress_cb=on_progress)
        excel_bytes = build_excel(results)
        out_path = RESULTS_DIR / f"{job_id}.xlsx"
        out_path.write_bytes(excel_bytes)

        # Annotate each result with its detected provider for the UI
        for r in results:
            r["provider"] = extract_provider(r["nameservers"])
        job["results"] = results
        job["status"] = "done"
    except Exception as e:
        job["status"] = "error"
        job["error"] = str(e)


# ---------------------------------------------------------------------------
# Routes
# ---------------------------------------------------------------------------

@app.post("/api/upload")
async def upload_csv(file: UploadFile = File(...)):
    if not file.filename.lower().endswith(".csv"):
        raise HTTPException(400, "Please upload a .csv file")

    content = await file.read()
    domains = parse_csv_domains(content)
    if not domains:
        raise HTTPException(400, "No domains found in column A of the CSV")

    job_id = str(uuid.uuid4())
    jobs[job_id] = {
        "status": "queued",
        "progress": 0,
        "total": len(domains),
        "results": [],
        "error": None,
    }

    thread = threading.Thread(target=run_job, args=(job_id, domains), daemon=True)
    thread.start()

    return {"job_id": job_id, "total": len(domains)}


@app.get("/api/status/{job_id}")
async def stream_status(job_id: str):
    """Server-Sent Events stream: emits progress and final results."""
    if job_id not in jobs:
        raise HTTPException(404, "Job not found")

    async def event_stream() -> AsyncGenerator[str, None]:
        while True:
            job = jobs.get(job_id)
            if not job:
                break

            data = {
                "status": job["status"],
                "progress": job["progress"],
                "total": job["total"],
            }

            if job["status"] == "done":
                data["results"] = job["results"]
                yield f"data: {json.dumps(data)}\n\n"
                break
            elif job["status"] == "error":
                data["error"] = job.get("error", "Unknown error")
                yield f"data: {json.dumps(data)}\n\n"
                break
            else:
                yield f"data: {json.dumps(data)}\n\n"

            await asyncio.sleep(0.4)

    return StreamingResponse(event_stream(), media_type="text/event-stream")


class ExportRequest(BaseModel):
    results: list[dict]
    starred_providers: list[str] = []


@app.post("/api/export")
async def export_filtered(req: ExportRequest):
    """Export visible results; splits into two sheets when starred providers are set."""
    excel_bytes = build_excel(req.results, starred_providers=req.starred_providers or None)
    filename = "ns_results_migration.xlsx" if req.starred_providers else "ns_results.xlsx"
    return Response(
        content=excel_bytes,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@app.get("/api/download/{job_id}")
async def download_excel(job_id: str):
    out_path = RESULTS_DIR / f"{job_id}.xlsx"
    if not out_path.exists():
        raise HTTPException(404, "Result file not found — job may still be running")
    return FileResponse(
        path=str(out_path),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="ns_results.xlsx",
    )


# Serve the frontend static files (mounted last so API routes take priority)
app.mount("/", StaticFiles(directory="/app/frontend", html=True), name="static")
